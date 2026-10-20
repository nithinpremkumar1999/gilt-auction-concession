"""
Stage 4 gate: a hand-check workbook for five auctions (one short, one medium, one long, one new
issue, one from Sept-Oct 2022) with every intermediate value of M1, M1b and M2 shown.

Inputs taken from the build (blue cells) are the values a spreadsheet cannot produce itself:
cash flows from `pricing.cash_flows()`, curve values z(t) from the Stage 2 spline, model yields
from the curve pricer, and DMO yields. Everything else is a live Excel formula (black): t, the
discount factors, present values, dirty and clean prices, the price at the model yield (which
must equal the curve dirty price), M1, the M1b residuals and baseline, the daily changes, the
control average, beta (SLOPE), abnormal changes, the cumulative path and the PM's three numbers.
Each sheet ends with the Python result for comparison.

Run:  python -m gilt_study.exercises.handcheck_stage4
Output: data/processed_data/measures/handcheck_5_auctions.xlsx (then recalculated with LibreOffice)
"""
import datetime as dt

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .. import auctions as au, measures as ms
from ..curve import curve_on, effective_settlement, load_panel, maturity_columns
from ..pricing import cash_flows

OUT_FILE = ms.OUT_DIR / "handcheck_5_auctions.xlsx"

# The five auctions (chosen at the Stage 4 build; see notebook 04 §5 for why each)
CHOSEN = {
    "short": None, "medium": None, "long": None, "new issue": None, "Sept-Oct 2022": None,
}

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
BOLD = Font(name=FONT, bold=True)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2A78D6")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")


def _hdr(ws, row, labels, col=1):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _put(ws, row, col, value, font=BLACK, fmt=None, fill=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def choose(M: pd.DataFrame, P: pd.DataFrame) -> dict:
    """Pick the five auctions: typical (beta inside 0-3, no rule-based first coupon), with M1b where
    the measure exists. Fixed rules so the choice is reproducible."""
    J = M.merge(P[["auction_date", "isin", "first_coupon_type", "first_coupon_source", "bid_to_cover"]],
                on=["auction_date", "isin"])
    ok = ~J["beta_outlier"] & ~J["flag_first_coupon_rule"] & ~J["drop_extrapolation_robustness"]
    pick = {}
    for b in ["short", "medium", "long"]:
        c = J[ok & (J["bucket"] == b) & J["m1b_bp"].notna() & ~J["new_issue"] & (J["years_beyond_curve"] == 0)]
        c = c[(pd.to_datetime(c["auction_date"]).dt.year >= 2010) & (pd.to_datetime(c["auction_date"]).dt.year <= 2016)]
        # the auction whose M1 is closest to its bucket's median (a typical case, not an extreme)
        med = c["m1_bp"].median()
        pick[b] = c.iloc[(c["m1_bp"] - med).abs().argsort().iloc[0]]["auction_id"]
    n = J[ok & J["new_issue"] & (J["first_coupon_type"] == "long")
          & J["first_coupon_source"].str.startswith("DMO")]
    n = n[(pd.to_datetime(n["auction_date"]).dt.year >= 2010)]
    pick["new issue"] = n.iloc[(n["m1_bp"] - n["m1_bp"].median()).abs().argsort().iloc[0]]["auction_id"]
    pick["Sept-Oct 2022"] = J[(J["auction_date"] == dt.date(2022, 10, 4))]["auction_id"].iloc[0]
    return pick


def m1_table(r, cp) -> tuple[pd.DataFrame, dict]:
    """The M1 calculation for one auction row (any object with the panel/measures fields), as a
    table of cash flows and a dict of totals - the same numbers as the workbook's M1 sheet."""
    crv, used, fb = curve_on(cp, r.t_minus_1)
    eff = effective_settlement(r.settlement_date, r.first_issue)
    cf = cash_flows(eff, r.maturity, r.coupon, r.first_issue, r.first_coupon)
    t = np.array([(d - eff).days / 365.25 for d in cf["dates"]])
    z = np.asarray(crv(t), float)
    T = pd.DataFrame({"cash_flow_date": cf["dates"], "amount": cf["amounts"], "t_years": t, "z_pct": z,
                      "beyond_curve": t > crv.t_max})
    T["df"] = np.exp(-T["z_pct"] / 100 * T["t_years"])
    T["pv"] = T["amount"] * T["df"]
    T["expo"] = cf["expo"]
    v = 1 / (1 + r.model_yield_t1 / 200)
    T["pv_at_model_yield"] = T["amount"] * v ** T["expo"]
    dirty = T["pv"].sum()
    out = dict(curve_date=used, effective_settlement=eff, t_max=crv.t_max, z_tmax=crv.z_tmax, f_tmax=crv.f_tmax,
               dirty=dirty, accrued=float(cf["accrued"]), clean=dirty - float(cf["accrued"]),
               model_yield=r.model_yield_t1, price_at_model_yield=T["pv_at_model_yield"].sum(),
               aay=r.aay, m1_bp=100 * (r.aay - r.model_yield_t1))
    return T, out


def _sheet_m1(wb, tag, r, cp):
    """M1: cash flows off the T-1 curve, discount factors, dirty price, model yield, M1."""
    ws = wb.create_sheet(f"{tag} M1")
    crv, used, fb = curve_on(cp, r.t_minus_1)
    eff = effective_settlement(r.settlement_date, r.first_issue)
    cf = cash_flows(eff, r.maturity, r.coupon, r.first_issue, r.first_coupon)
    ws["A1"] = f"{tag}: {r.gilt_name}, auction {r.auction_date:%d %b %Y} - M1 (spec §7)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    inputs = [("Auction date", r.auction_date, "dd-mmm-yyyy"), ("Settlement date (DMO issue date)", r.settlement_date, "dd-mmm-yyyy"),
              ("Effective settlement (first issue date if later, decision #10)", eff, "dd-mmm-yyyy"),
              ("Curve date (T-1 close)", used, "dd-mmm-yyyy"), ("Coupon (% p.a.)", r.coupon, "0.000"),
              ("Maturity", r.maturity, "dd-mmm-yyyy"), ("First issue date", r.first_issue, "dd-mmm-yyyy"),
              ("First coupon date (source below)", r.first_coupon, "dd-mmm-yyyy"),
              ("First coupon source", r.first_coupon_source, None),
              ("Accrued interest at settlement (cash_flows)", float(cf["accrued"]), "0.000000"),
              ("Average accepted yield, AAY (% , D2.1A)", r.aay, "0.000000"),
              ("Curve: last maturity T (y)", crv.t_max, "0.0"), ("Curve: z(T) (%)", crv.z_tmax, "0.000000"),
              ("Curve: BoE forward f(T) (%)", crv.f_tmax, "0.000000")]
    for i, (lab, val, fmt) in enumerate(inputs):
        _put(ws, 3 + i, 1, lab, BOLD)
        _put(ws, 3 + i, 2, val, BLUE, fmt)
    R_SETTLE, R_ACC, R_AAY = 5, 12, 13          # rows of effective settlement, accrued, AAY
    top = 3 + len(inputs) + 2
    _hdr(ws, top, ["Cash flow date", "Amount per £100", "t = days/365.25", "z(t) % (spline; flat fwd beyond T)",
                   "DF = exp(-z/100·t)", "PV = amount·DF", "Exponent (semi-annual periods)", "v^exponent at model yield",
                   "PV at model yield", "BoE point below", "z below", "BoE point above", "z above",
                   "Linear interp. of z (check)"])
    knots_t, knots_z = crv.t, crv.zs
    rows = []
    for j, (d, a, e) in enumerate(zip(cf["dates"], cf["amounts"], cf["expo"])):
        rr = top + 1 + j
        t = (d - eff).days / 365.25
        _put(ws, rr, 1, d, BLUE, "dd-mmm-yyyy")
        _put(ws, rr, 2, float(a), BLUE, "0.000000")
        _put(ws, rr, 3, f"=(A{rr}-$B${R_SETTLE})/365.25", BLACK, "0.000000")
        _put(ws, rr, 4, float(crv(t)), BLUE, "0.000000")
        _put(ws, rr, 5, f"=EXP(-D{rr}/100*C{rr})", BLACK, "0.00000000")
        _put(ws, rr, 6, f"=B{rr}*E{rr}", BLACK, "0.000000")
        _put(ws, rr, 7, float(e), BLUE, "0.000000")
        _put(ws, rr, 8, f"=(1/(1+$B${top + len(cf['dates']) + 5}/200))^G{rr}", BLACK, "0.00000000")
        _put(ws, rr, 9, f"=B{rr}*H{rr}", BLACK, "0.000000")
        lo = knots_t[knots_t <= t]
        hi = knots_t[knots_t >= t]
        if len(lo) and len(hi):
            tl, th = lo[-1], hi[0]
            zl, zh = knots_z[knots_t == tl][0], knots_z[knots_t == th][0]
            _put(ws, rr, 10, float(tl), BLUE, "0.000"); _put(ws, rr, 11, float(zl), BLUE, "0.000000")
            _put(ws, rr, 12, float(th), BLUE, "0.000"); _put(ws, rr, 13, float(zh), BLUE, "0.000000")
            _put(ws, rr, 14, f"=IF(L{rr}=J{rr},K{rr},K{rr}+(M{rr}-K{rr})*(C{rr}-J{rr})/(L{rr}-J{rr}))", BLACK, "0.000000")
        rows.append(rr)
    last = rows[-1]
    s = last + 2
    out = [("Curve dirty price = SUM(PV)", f"=SUM(F{rows[0]}:F{last})", "0.000000"),
           ("Accrued interest", f"=B{R_ACC}", "0.000000"),
           ("Curve clean price", f"=B{s}-B{s + 1}", "0.000000"),
           ("Model yield (% , brentq in Python)", None, "0.00000000"),
           ("Price at the model yield = SUM(PV at model yield)", f"=SUM(I{rows[0]}:I{last})", "0.000000"),
           ("Check: price at model yield - curve dirty (≈ 0)", f"=B{s + 4}-B{s}", "0.000E+00"),
           ("M1 (bp) = 100·(AAY - model yield)", f"=100*(B{R_AAY}-B{s + 3})", "0.000"),
           ("Python M1 (bp), for comparison", float(r.m1_bp), "0.000")]
    for i, (lab, val, fmt) in enumerate(out):
        _put(ws, s + i, 1, lab, BOLD)
        if val is None:
            _put(ws, s + i, 2, float(r.model_yield_t1), BLUE, fmt)
        else:
            _put(ws, s + i, 2, val, BLUE if isinstance(val, float) else BLACK, fmt,
                 KEY_FILL if lab.startswith("M1") else None)
    ws.cell(row=s + 3, column=2).comment = Comment("Solved in Python: the semi-annual GRY at which gilt_price() "
                                                    "reproduces the curve dirty price (brentq). The two rows below "
                                                    "re-price it here with the GRY formula to show it does.", "build")
    # fix the model-yield reference used in column H (row s + 3)
    for rr in rows:
        ws.cell(row=rr, column=8, value=f"=(1/(1+$B${s + 3}/200))^G{rr}")
    ws.column_dimensions["A"].width = 46
    for col in range(2, 15):
        ws.column_dimensions[get_column_letter(col)].width = 15
    return ws, f"'{tag} M1'!B{s + 6}"


def _sheet_m1b(wb, tag, r, resid, m1_ref):
    ws = wb.create_sheet(f"{tag} M1b")
    ws["A1"] = f"{tag}: {r.gilt_name}, auction {r.auction_date:%d %b %Y} - M1b baseline, T-25 to T-6 (decision #4)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    _hdr(ws, 3, ["k (business days)", "Date", "DMO reference yield (%)", "Model yield (%, curve of that day)",
                 "Residual (bp) = 100·(DMO - model)", "Included (1/0)", "Excluded because"])
    R = resid[resid["auction_id"] == r.auction_id].sort_values("k")
    for j, x in enumerate(R.itertuples(index=False)):
        rr = 4 + j
        _put(ws, rr, 1, int(x.k), BLUE)
        _put(ws, rr, 2, x.date, BLUE, "dd-mmm-yyyy")
        _put(ws, rr, 3, None if pd.isna(x.dmo_yield) else float(x.dmo_yield), BLUE, "0.000000")
        _put(ws, rr, 4, None if pd.isna(x.model_yield) else float(x.model_yield), BLUE, "0.000000")
        _put(ws, rr, 5, f'=IF(F{rr}=1,100*(C{rr}-D{rr}),"")', BLACK, "0.000")
        _put(ws, rr, 6, 0 if isinstance(x.excluded, str) else 1, BLUE)
        _put(ws, rr, 7, x.excluded if isinstance(x.excluded, str) else "", BLUE)
    last = 3 + len(R)
    s = last + 2
    out = [("Valid days (need ≥ 10)", f"=SUM(F4:F{last})", "0"),
           ("Baseline = mean residual (bp)", f"=IF(B{s}>=10,AVERAGEIF(F4:F{last},1,E4:E{last}),\"n/a\")", "0.000"),
           ("M1 (bp), from the M1 sheet", f"={m1_ref}", "0.000"),
           ("M1b (bp) = M1 - baseline", f"=IF(B{s}>=10,B{s + 2}-B{s + 1},\"n/a\")", "0.000"),
           ("Python M1b (bp), for comparison", None if pd.isna(r.m1b_bp) else float(r.m1b_bp), "0.000"),
           ("Python note", r.m1b_missing_reason if isinstance(r.m1b_missing_reason, str) else "", None)]
    for i, (lab, val, fmt) in enumerate(out):
        _put(ws, s + i, 1, lab, BOLD)
        _put(ws, s + i, 2, val, BLUE if (val is None or isinstance(val, float) or lab == "Python note") else BLACK, fmt,
             KEY_FILL if lab.startswith("M1b (bp)") else None)
    ws.column_dimensions["A"].width = 34
    for col, w in zip("BCDEFG", [13, 16, 18, 18, 10, 52]):
        ws.column_dimensions[col].width = w
    return ws


def _sheet_m2(wb, tag, r, cp, Y):
    """M2: model yields T-121..T+10, control maturities, beta, AR, CAR, the PM's three numbers."""
    days = ms.event_days(cp.index, r.auction_date, ms.FIRST_K, ms.M2_WINDOW[1])
    ks = list(range(ms.FIRST_K, ms.M2_WINDOW[1] + 1))
    mats = maturity_columns(cp)
    Z = cp.loc[days, mats]
    avail = [m for m in mats if Z[m].notna().all()]
    cm = list(ms.control_maturities(r.years_to_maturity, avail))
    ylook = Y.set_index(["isin", "curve_date"])["model_yield"]

    wz = wb.create_sheet(f"{tag} control")
    wz["A1"] = (f"{tag}: control maturities = BoE points 2-30y more than 3y from the gilt's "
                f"{r.years_to_maturity:.2f}y to maturity ({len(cm)} points). Top: z (%); bottom: daily change (bp).")
    wz["A1"].font = Font(name=FONT, bold=True, size=12)
    _hdr(wz, 3, ["k", "Date"] + [f"z {m:g}y" for m in cm])
    for j, (k, ts) in enumerate(zip(ks, days)):
        rr = 4 + j
        _put(wz, rr, 1, k, BLUE); _put(wz, rr, 2, ts.date(), BLUE, "dd-mmm-yyyy")
        for c, m in enumerate(cm):
            _put(wz, rr, 3 + c, float(Z.at[ts, m]), BLUE, "0.000000")
    top2 = 4 + len(ks) + 2
    ncol = 2 + len(cm)
    _hdr(wz, top2, ["k", "Date"] + [f"Δz {m:g}y (bp)" for m in cm] + ["Control Δc = average (bp)"])
    for j, k in enumerate(ks[1:], start=1):
        rr = top2 + j
        src, prev = 4 + j, 3 + j
        _put(wz, rr, 1, f"=A{src}", BLACK); _put(wz, rr, 2, f"=B{src}", BLACK, "dd-mmm-yyyy")
        for c in range(len(cm)):
            col = get_column_letter(3 + c)
            _put(wz, rr, 3 + c, f"=100*({col}{src}-{col}{prev})", BLACK, "0.0000")
        _put(wz, rr, ncol + 1, f"=AVERAGE(C{rr}:{get_column_letter(ncol)}{rr})", BLACK, "0.0000")
    dc_col = get_column_letter(ncol + 1)

    ws = wb.create_sheet(f"{tag} M2")
    ws["A1"] = f"{tag}: {r.gilt_name}, auction {r.auction_date:%d %b %Y} - M2 with control B-β (decision #13)"
    ws["A1"].font = Font(name=FONT, bold=True, size=12)
    _hdr(ws, 3, ["k", "Date", "Model yield y (%)", "Δy (bp)", "Δc (bp), from control sheet", "In β window (1/0)",
                 "AR B-β = Δy - β·Δc", "CAR B-β (base T-6)", "AR B (β = 1)", "CAR B", "Raw cumulative Δy"])
    first = 4
    rowk = {}
    for j, (k, ts) in enumerate(zip(ks, days)):
        rr = first + j
        rowk[k] = rr
        _put(ws, rr, 1, k, BLUE); _put(ws, rr, 2, ts.date(), BLUE, "dd-mmm-yyyy")
        _put(ws, rr, 3, float(ylook[(r.isin, ts.date())]), BLUE, "0.00000000")
        if j:
            _put(ws, rr, 4, f"=100*(C{rr}-C{rr - 1})", BLACK, "0.0000")
            _put(ws, rr, 5, f"='{tag} control'!{dc_col}{top2 + j}", Font(name=FONT, color="008000"), "0.0000")
        _put(ws, rr, 6, 1 if ms.BETA_WINDOW[0] <= k <= ms.BETA_WINDOW[1] else 0, BLUE)
    lastk = rowk[ms.M2_WINDOW[1]]
    s = lastk + 2
    b_lo, b_hi = rowk[ms.BETA_WINDOW[0]], rowk[ms.BETA_WINDOW[1]]
    stats = [("β = SLOPE(Δy, Δc) over k = -120..-21", f"=SLOPE(D{b_lo}:D{b_hi},E{b_lo}:E{b_hi})", "0.0000"),
             ("α (estimated, not subtracted)", f"=INTERCEPT(D{b_lo}:D{b_hi},E{b_lo}:E{b_hi})", "0.0000"),
             ("R²", f"=RSQ(D{b_lo}:D{b_hi},E{b_lo}:E{b_hi})", "0.000"),
             ("Observations", f"=COUNT(D{b_lo}:D{b_hi})", "0"),
             ("Python β, for comparison", float(r.beta), "0.0000")]
    for i, (lab, val, fmt) in enumerate(stats):
        _put(ws, s + i, 1, lab, BOLD)
        _put(ws, s + i, 2, val, BLUE if isinstance(val, float) else BLACK, fmt, KEY_FILL if i == 0 else None)
    BETA = f"$B${s}"
    base = rowk[ms.M2_BASE]
    for k in range(ms.M2_WINDOW[0], ms.M2_WINDOW[1] + 1):
        rr = rowk[k]
        _put(ws, rr, 7, f"=D{rr}-{BETA}*E{rr}", BLACK, "0.0000")
        _put(ws, rr, 9, f"=D{rr}-E{rr}", BLACK, "0.0000")
        if k == ms.M2_BASE:
            for col in (8, 10, 11):
                _put(ws, rr, col, 0, BLACK, "0.0000")
        elif k > ms.M2_BASE:
            _put(ws, rr, 8, f"=H{rr - 1}+G{rr}", BLACK, "0.0000")
            _put(ws, rr, 10, f"=J{rr - 1}+I{rr}", BLACK, "0.0000")
            _put(ws, rr, 11, f"=K{rr - 1}+D{rr}", BLACK, "0.0000")
        else:
            _put(ws, rr, 8, f"=H{rr + 1}-G{rr + 1}", BLACK, "0.0000")
            _put(ws, rr, 10, f"=J{rr + 1}-I{rr + 1}", BLACK, "0.0000")
            _put(ws, rr, 11, f"=K{rr + 1}-D{rr + 1}", BLACK, "0.0000")
    p = s + len(stats) + 1
    _hdr(ws, p, ["The PM's three numbers (bp)", "B-β (main)", "B (β = 1)", "Raw (no control)", "Python B-β"])
    for i, (label, (a, b)) in enumerate(ms.PM_NUMBERS.items()):
        rr = p + 1 + i
        name = {"pre": "Pre-auction cheapening, CAR(T-1) - CAR(T-6)", "day": "Auction day, CAR(T) - CAR(T-1)",
                "post": "Post-auction, CAR(T+5) - CAR(T)"}[label]
        _put(ws, rr, 1, name, BOLD)
        _put(ws, rr, 2, f"=H{rowk[b]}-H{rowk[a]}", BLACK, "0.000", KEY_FILL)
        _put(ws, rr, 3, f"=J{rowk[b]}-J{rowk[a]}", BLACK, "0.000")
        _put(ws, rr, 4, f"=K{rowk[b]}-K{rowk[a]}", BLACK, "0.000")
        _put(ws, rr, 5, float(getattr(r, f"m2_{label}_bbeta_bp")), BLUE, "0.000")
    ws.column_dimensions["A"].width = 44
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 14
    ws.freeze_panes = "C4"
    wz.freeze_panes = "C4"
    return ws


def build(write: bool = True):
    P = au.load_auction_panel()
    M = ms.load_measures()
    cp = load_panel()
    Y = ms.load_model_yields()
    resid = pd.read_csv(ms.RESIDUALS_FILE, float_precision="round_trip")
    resid["date"] = [dt.date.fromisoformat(v) for v in resid["date"]]
    picks = choose(M, P)
    J = M.merge(P[["auction_date", "isin", "coupon", "maturity", "first_issue", "first_coupon", "first_coupon_source"]],
                on=["auction_date", "isin"])
    J["first_issue"] = [v if isinstance(v, dt.date) else None for v in J["first_issue"]]
    J["first_coupon"] = [v if isinstance(v, dt.date) else None for v in J["first_coupon"]]

    wb = Workbook()
    wr = wb.active
    wr.title = "README"
    wr["A1"] = "Stage 4 hand check: five auctions, every intermediate value of M1, M1b and M2"
    wr["A1"].font = Font(name=FONT, bold=True, size=13)
    notes = ["Blue cells: inputs from the build (cash flows, spline z(t), model yields, DMO yields, BoE curve points).",
             "Black cells: live formulas. Green cells: links to another sheet. Yellow: the results to check.",
             "Each sheet ends with the Python result next to the spreadsheet's own, so the two can be compared.",
             "Definitions: build spec §7 and src/gilt_study/measures.py. Event days k are UK business days (= BoE curve dates).",
             "M1 = 100·(AAY - model yield off the T-1 curve, for the auction settlement date).",
             "M1b = M1 - mean residual over T-25..T-6 (days within 5 business days of an earlier auction of the gilt excluded; ≥ 10 days).",
             "M2: AR = Δy - β·Δc with β = SLOPE over k = -120..-21; CAR = 0 at T-6."]
    for i, n in enumerate(notes):
        _put(wr, 3 + i, 1, n, Font(name=FONT))
    _hdr(wr, 11, ["Slot", "Auction", "Gilt", "Bucket", "Regime", "New issue", "Sheets"])
    tags = {"short": "A1", "medium": "A2", "long": "A3", "new issue": "A4", "Sept-Oct 2022": "A5"}
    for i, (slot, aid) in enumerate(picks.items()):
        r = next(J[J["auction_id"] == aid].itertuples(index=False))
        tag = tags[slot]
        _, m1_ref = _sheet_m1(wb, tag, r, cp)
        _sheet_m1b(wb, tag, r, resid, m1_ref)
        _sheet_m2(wb, tag, r, cp, Y)
        row = [slot, r.auction_date, r.gilt_name, r.bucket, r.regime, bool(r.new_issue), f"{tag} M1 / M1b / M2 / control"]
        for j, v in enumerate(row):
            _put(wr, 12 + i, 1 + j, v, Font(name=FONT), "dd-mmm-yyyy" if j == 1 else None)
    for col, w in zip("ABCDEFG", [16, 13, 32, 9, 12, 10, 30]):
        wr.column_dimensions[col].width = w
    wr.column_dimensions["A"].width = 16
    if write:
        ms.OUT_DIR.mkdir(parents=True, exist_ok=True)
        wb.save(OUT_FILE)
    return picks


if __name__ == "__main__":
    print(build())
